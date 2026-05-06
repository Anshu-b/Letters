import csv
from docx import Document
import os

# Path to the CSV file
csv_file = 'judges.csv'

# Path to the template DOCX file (you need to provide this)
template_file = 'template.docx'

# Output directory for generated letters
output_dir = 'generated letters'

# Check if template exists
if not os.path.exists(template_file):
    print(f"Template file '{template_file}' not found. Please provide the .docx template file.")
    exit(1)

# Ensure output folder exists
os.makedirs(output_dir, exist_ok=True)

from docx.shared import Pt

# Read the CSV
with open(csv_file, 'r') as f:
    reader = csv.DictReader(f)
    for row in reader:
        name = row['Name'].strip()
        track = row['Track'].strip()
        first_name = row['First Name'].strip()
        
        # Load the template
        doc = Document(template_file)
        
        # Function to replace text in paragraphs
        def replace_text_in_paragraphs(paragraphs, old_text, new_text):
            for para in paragraphs:
                if old_text in para.text:
                    para.text = para.text.replace(old_text, new_text)
        
        # Function to replace text in tables
        def replace_text_in_tables(tables, old_text, new_text):
            for table in tables:
                for row in table.rows:
                    for cell in row.cells:
                        replace_text_in_paragraphs(cell.paragraphs, old_text, new_text)
                        replace_text_in_tables(cell.tables, old_text, new_text)
        
        # Function to enforce font size on paragraphs
        def enforce_font_size_in_paragraphs(paragraphs, size=11):
            for para in paragraphs:
                for run in para.runs:
                    run.font.size = Pt(size)
        
        # Function to enforce font size in tables
        def enforce_font_size_in_tables(tables, size=11):
            for table in tables:
                for row in table.rows:
                    for cell in row.cells:
                        enforce_font_size_in_paragraphs(cell.paragraphs, size)
                        enforce_font_size_in_tables(cell.tables, size)
        
        # Replace {{Name}}, {{Track}}, and {{First Name}}
        replace_text_in_paragraphs(doc.paragraphs, '{{Name}}', name)
        replace_text_in_paragraphs(doc.paragraphs, '{{Track}}', track)
        replace_text_in_paragraphs(doc.paragraphs, '{{First Name}}', first_name)
        replace_text_in_tables(doc.tables, '{{Name}}', name)
        replace_text_in_tables(doc.tables, '{{Track}}', track)
        replace_text_in_tables(doc.tables, '{{First Name}}', first_name)
        
        # Enforce 11pt formatting in body, headers, and footers
        enforce_font_size_in_paragraphs(doc.paragraphs, 11)
        enforce_font_size_in_tables(doc.tables, 11)
        for section in doc.sections:
            enforce_font_size_in_paragraphs(section.header.paragraphs, 11)
            enforce_font_size_in_tables(section.header.tables, 11)
            enforce_font_size_in_paragraphs(section.footer.paragraphs, 11)
            enforce_font_size_in_tables(section.footer.tables, 11)
        
        # Save the personalized document
        output_file = os.path.join(output_dir, f"{name.replace(' ', '_')}_letter.docx")
        doc.save(output_file)
        print(f"Generated: {output_file}")